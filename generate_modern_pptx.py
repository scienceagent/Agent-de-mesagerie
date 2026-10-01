import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Colors
    c_bg = RGBColor(248, 250, 252)        # #f8fafc
    c_primary = RGBColor(37, 99, 235)     # #2563eb royal blue
    c_deep_blue = RGBColor(29, 78, 216)   # #1d4ed8
    c_soft_blue = RGBColor(239, 246, 255) # #eff6ff
    c_dark = RGBColor(15, 23, 42)         # #0f172a
    c_muted = RGBColor(100, 116, 139)     # #64748b
    c_white = RGBColor(255, 255, 255)
    c_border = RGBColor(226, 232, 240)
    c_green = RGBColor(22, 163, 74)

    def add_base_slide(header_text="FCIM • SISTEME DISTRIBUITE 2026", step_text="Etapa 1"):
        slide = prs.slides.add_slide(blank_layout)
        # Background shape
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = c_bg
        bg.line.fill.background()

        # Header tag
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8), Inches(0.4))
        p = tb.text_frame.paragraphs[0]
        p.text = f"●  {header_text}"
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = c_primary

        # Right badge
        tb2 = slide.shapes.add_textbox(Inches(9.5), Inches(0.4), Inches(3), Inches(0.4))
        p2 = tb2.text_frame.paragraphs[0]
        p2.alignment = PP_ALIGN.RIGHT
        p2.text = step_text
        p2.font.name = "Arial"
        p2.font.size = Pt(10)
        p2.font.bold = True
        p2.font.color.rgb = c_muted

        return slide

    def add_titles(slide, eyebrow, title, desc=None):
        # Eyebrow
        tb_eye = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(10), Inches(0.35))
        p_eye = tb_eye.text_frame.paragraphs[0]
        p_eye.text = eyebrow.upper()
        p_eye.font.name = "Arial"
        p_eye.font.size = Pt(11)
        p_eye.font.bold = True
        p_eye.font.color.rgb = c_primary

        # Title
        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(11.5), Inches(0.8))
        p_t = tb_title.text_frame.paragraphs[0]
        p_t.text = title
        p_t.font.name = "Arial"
        p_t.font.size = Pt(28)
        p_t.font.bold = True
        p_t.font.color.rgb = c_dark

        if desc:
            tb_desc = slide.shapes.add_textbox(Inches(0.8), Inches(1.85), Inches(11.5), Inches(0.5))
            p_d = tb_desc.text_frame.paragraphs[0]
            p_d.text = desc
            p_d.font.name = "Arial"
            p_d.font.size = Pt(13)
            p_d.font.color.rgb = c_muted

    # ==================== SLIDE 1: COVER ====================
    s1 = add_base_slide(header_text="FCIM • INGINERIA SISTEMELOR SOFTWARE 2026", step_text="PROIECT DE TEZĂ")
    
    # Title Box
    tb_c = s1.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.5), Inches(2.2))
    tf_c = tb_c.text_frame
    p_c1 = tf_c.paragraphs[0]
    p_c1.text = "PROIECTAREA ȘI IMPLEMENTAREA UNUI"
    p_c1.font.name = "Arial"
    p_c1.font.size = Pt(18)
    p_c1.font.bold = True
    p_c1.font.color.rgb = c_primary
    
    p_c2 = tf_c.add_paragraph()
    p_c2.text = "Agent de Mesagerie Asincron Rezilient"
    p_c2.font.name = "Arial"
    p_c2.font.size = Pt(36)
    p_c2.font.bold = True
    p_c2.font.color.rgb = c_dark

    p_c3 = tf_c.add_paragraph()
    p_c3.text = "Garanții At-Least-Once, Deduplicare Idempotentă, Validare XSD și Rutare Unicast/Multicast"
    p_c3.font.name = "Arial"
    p_c3.font.size = Pt(14)
    p_c3.font.color.rgb = c_muted

    # Team Members row
    team_members = [
        ("Raevschi Grigore", "Core Broker & Concurrency"),
        ("Bajureanu Patricia", "Contract XSD & Validare"),
        ("Tihon Stanislav", "Protocol ACK/NACK & Retry"),
        ("Cojocaru Mihai", "TCP Framing & Queues")
    ]
    card_w = Inches(2.7)
    gap = Inches(0.25)
    left_start = Inches(0.8)

    for i, (name, role) in enumerate(team_members):
        cur_left = left_start + i * (card_w + gap)
        card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cur_left, Inches(4.5), card_w, Inches(1.8))
        card.fill.solid()
        card.fill.fore_color.rgb = c_white
        card.line.color.rgb = c_border
        card.line.width = Pt(1)

        # Text in card
        tf_tm = card.text_frame
        tf_tm.margin_left = Inches(0.2)
        tf_tm.margin_top = Inches(0.3)
        p_n = tf_tm.paragraphs[0]
        p_n.text = name
        p_n.font.name = "Arial"
        p_n.font.size = Pt(14)
        p_n.font.bold = True
        p_n.font.color.rgb = c_dark

        p_r = tf_tm.add_paragraph()
        p_r.text = role
        p_r.font.name = "Arial"
        p_r.font.size = Pt(11)
        p_r.font.color.rgb = c_muted

    # ==================== SLIDE 2: PROBLEMA ====================
    s2 = add_base_slide()
    add_titles(s2, "01. Context & Provocări", "De ce eșuează sistemele distribuite nominale?")
    
    problems = [
        ("01", "Pierderea Mesajelor", "În HTTP/REST sincron, dacă un consumator devine indisponibil 200ms, tranzacția eșuează imediat cu 503, provocând pierderi de comenzi critice.", False),
        ("02", "Crash Post-Procesare", "Problema celor Doi Generali: dacă serverul aplică debitarea locală dar cade înainte de ACK, mesajul retransmis provoacă duplicate nedorite!", True),
        ("03", "Mesaje Malformate", "Lipsa validării contractului formal (XML XSD) permite mesaje de tip Poison Pill care blochează cozile și blochează procesarea.", False)
    ]
    col_w = Inches(3.68)
    for i, (num, title, body, is_blue) in enumerate(problems):
        c_x = Inches(0.8) + i * (col_w + Inches(0.35))
        box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x, Inches(2.4), col_w, Inches(4.2))
        box.fill.solid()
        box.fill.fore_color.rgb = c_primary if is_blue else c_white
        box.line.color.rgb = c_primary if is_blue else c_border
        
        tf = box.text_frame
        tf.margin_left = Inches(0.3)
        tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.4)
        
        p1 = tf.paragraphs[0]
        p1.text = num
        p1.font.size = Pt(24)
        p1.font.bold = True
        p1.font.color.rgb = c_white if is_blue else c_primary
        
        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(18)
        p2.font.bold = True
        p2.font.color.rgb = c_white if is_blue else c_dark

        p3 = tf.add_paragraph()
        p3.text = "\n" + body
        p3.font.size = Pt(13)
        p3.font.color.rgb = c_white if is_blue else c_muted

    # ==================== SLIDE 3: VALOARE SI METRICI ====================
    s3 = add_base_slide()
    add_titles(s3, "02. Valoare Adăugată", "Metrici Cheie de Reziliență și Garanții")
    
    # Left big stat card
    b1 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.4), Inches(5.6), Inches(4.2))
    b1.fill.solid()
    b1.fill.fore_color.rgb = c_soft_blue
    b1.line.color.rgb = c_primary
    tf1 = b1.text_frame
    tf1.margin_left = Inches(0.4)
    tf1.margin_top = Inches(0.4)
    
    p = tf1.paragraphs[0]
    p.text = "100%"
    p.font.size = Pt(64)
    p.font.bold = True
    p.font.color.rgb = c_primary
    
    p = tf1.add_paragraph()
    p.text = "RATĂ DE SUCCES LA TESTELE DE HAOS"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = c_muted
    
    p = tf1.add_paragraph()
    p.text = "\nToate cele 6 suite automate de verificare (XML validation, ordering monotonic, ACK timeout, deduplicare idempotentă, crash recovery și unicast queues) rulează cu succes pe codul actual."
    p.font.size = Pt(14)
    p.font.color.rgb = c_dark

    # Right stats
    right_stats = [
        ("0", "Pierderi de Date", "Jurnal persistent JSONL pe disc"),
        ("4x", "Thread Pool Paralel", "Dispatch concurent non-blocant"),
    ]
    for i, (num, lbl, sub) in enumerate(right_stats):
        bx = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.4) + i * Inches(1.5), Inches(5.7), Inches(1.3))
        bx.fill.solid()
        bx.fill.fore_color.rgb = c_white
        bx.line.color.rgb = c_border
        tf = bx.text_frame
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.2)
        p = tf.paragraphs[0]
        p.text = f"{num}   {lbl.upper()}"
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = c_primary
        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.size = Pt(12)
        p2.font.color.rgb = c_muted

    # Formula banner below
    bf = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(5.5), Inches(5.7), Inches(1.1))
    bf.fill.solid()
    bf.fill.fore_color.rgb = c_deep_blue
    bf.line.fill.background()
    tff = bf.text_frame
    tff.margin_left = Inches(0.3)
    tff.margin_top = Inches(0.2)
    p = tff.paragraphs[0]
    p.text = "FORMULA DE REZILIENȚĂ:"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = c_white
    p2 = tff.add_paragraph()
    p2.text = "At-Least-Once + Idempotent = Exactly-Once la nivel de business"
    p2.font.size = Pt(13)
    p2.font.color.rgb = c_white

    # ==================== SLIDE 4: TOPOLOGIE ARHITECTURALA ====================
    s4 = add_base_slide()
    add_titles(s4, "03. Arhitectură & Topologie", "Topologie Decuplată pe Trei Niveluri")

    tiers = [
        ("Producător", "Publisher (Python / .NET)", "Generează comenzi structurate, semnează sarcina cu identificator unic UUID și folosește delimitare strictă prin terminator newline (\\n).", False),
        ("Nucleu Central", "Broker (.NET 8 Core)", "Găzduiește socketul TCP asincron, pipeline-ul de validare XSD, secvențiatorul topic-urilor, registrul de conexiuni și motorul de dispatching pe 4 fire de execuție.", True),
        ("Consumator", "Subscriber Idempotent", "Consumă mesaje, menține jurnalul ID-urilor procesate în memorie/disc și emite ACK. Ignoră duplicatul dacă a fost deja executat.", False)
    ]
    for i, (badge, title, desc, is_blue) in enumerate(tiers):
        c_x = Inches(0.8) + i * (col_w + Inches(0.35))
        box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x, Inches(2.4), col_w, Inches(4.2))
        box.fill.solid()
        box.fill.fore_color.rgb = c_primary if is_blue else c_white
        box.line.color.rgb = c_primary if is_blue else c_border
        
        tf = box.text_frame
        tf.margin_left = Inches(0.3)
        tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.4)
        
        p1 = tf.paragraphs[0]
        p1.text = badge.upper()
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = c_white if is_blue else c_primary
        
        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(18)
        p2.font.bold = True
        p2.font.color.rgb = c_white if is_blue else c_dark

        p3 = tf.add_paragraph()
        p3.text = "\n" + desc
        p3.font.size = Pt(13)
        p3.font.color.rgb = c_white if is_blue else c_muted

    # ==================== SLIDE 5: XSD & DLQ ====================
    s5 = add_base_slide()
    add_titles(s5, "04. Contract de Mesaj", "Validare Formală XSD & Dead-Letter Queue (DLQ)")

    b_x1 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.4), Inches(5.6), Inches(4.2))
    b_x1.fill.solid()
    b_x1.fill.fore_color.rgb = c_white
    b_x1.line.color.rgb = c_border
    tfx1 = b_x1.text_frame
    tfx1.margin_left = Inches(0.3)
    tfx1.margin_top = Inches(0.4)
    p = tfx1.paragraphs[0]
    p.text = "VALIDARE LA FRONTIERĂ (FAIL-FAST)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = c_primary
    p2 = tfx1.add_paragraph()
    p2.text = "\n• Orice mesaj FORMAT:XML este verificat împotriva schemei message.xsd înainte de acceptare.\n• Respinge tag-uri lipsă sau tipuri de date incompatibile.\n• Previne atacurile de tip XML Injection și coruperea memoriei.\n• Elimină riscul ca un payload malformat să ajungă la consumatori."
    p2.font.size = Pt(13)
    p2.font.color.rgb = c_muted

    b_x2 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.4), Inches(5.7), Inches(4.2))
    b_x2.fill.solid()
    b_x2.fill.fore_color.rgb = c_soft_blue
    b_x2.line.color.rgb = c_primary
    tfx2 = b_x2.text_frame
    tfx2.margin_left = Inches(0.3)
    tfx2.margin_top = Inches(0.4)
    p = tfx2.paragraphs[0]
    p.text = "DEAD-LETTER QUEUE (DLQ)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = c_deep_blue
    p2 = tfx2.add_paragraph()
    p2.text = "\n• Mesajele malformate sau cele cu retries epuizate sunt deviate automat în storage/dead_letter.journal.\n• Coada activă nu este niciodată blocată de un Poison Pill.\n• Permite auditarea și re-procesarea manuală după remedierea erorilor."
    p2.font.size = Pt(13)
    p2.font.color.rgb = c_dark

    # ==================== SLIDE 6: ACK/NACK ====================
    s6 = add_base_slide()
    add_titles(s6, "05. Protocol Bidirecțional", "Mecanismul ACK, NACK și Retransmisie")

    steps_ack = [
        ("01", "In-Flight State", "Mesajul transmis intră în registrul InFlightMessages asociat cu un timestamp de expirare (ACK Timeout)."),
        ("02", "Confirmare (ACK)", "Clientul confirmă cu ACK <uuid>. Brokerul șterge mesajul din zbor și finalizează livrarea cu succes."),
        ("03", "NACK / Timeout", "Dacă se primește NACK sau expiră timeout-ul, brokerul retransmite mesajul automat către următorul consumator disponibil.")
    ]
    for i, (num, title, body) in enumerate(steps_ack):
        c_x = Inches(0.8) + i * (col_w + Inches(0.35))
        box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x, Inches(2.4), col_w, Inches(4.2))
        box.fill.solid()
        box.fill.fore_color.rgb = c_primary if i == 1 else c_white
        box.line.color.rgb = c_primary if i == 1 else c_border
        tf = box.text_frame
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.4)
        p1 = tf.paragraphs[0]
        p1.text = num
        p1.font.size = Pt(24)
        p1.font.bold = True
        p1.font.color.rgb = c_white if i == 1 else c_primary
        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(18)
        p2.font.bold = True
        p2.font.color.rgb = c_white if i == 1 else c_dark
        p3 = tf.add_paragraph()
        p3.text = "\n" + body
        p3.font.size = Pt(13)
        p3.font.color.rgb = c_white if i == 1 else c_muted

    # ==================== SLIDE 7: SCENARIUL CRITIC ====================
    s7 = add_base_slide()
    add_titles(s7, "06. Verificare Inginerească", "Scenariul Critic: Crash Post-Procesare")

    bc1 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.4), Inches(5.6), Inches(4.2))
    bc1.fill.solid()
    bc1.fill.fore_color.rgb = c_white
    bc1.line.color.rgb = c_border
    tfc1 = bc1.text_frame
    tfc1.margin_left = Inches(0.3)
    tfc1.margin_top = Inches(0.4)
    p = tfc1.paragraphs[0]
    p.text = "DESFĂȘURAREA SCENARIULUI"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = c_primary
    p2 = tfc1.add_paragraph()
    p2.text = "\n1. Producătorul trimite comanda de debitare.\n2. Consumatorul aplică efectul local în baza de date.\n3. Consumatorul se prăbușește brusc înainte de ACK!\n4. Brokerul detectează deconectarea și retransmite mesajul.\n5. Consumatorul repornește, detectează ID-ul în istoricul local și trimite ACK fără a repeta tranzacția."
    p2.font.size = Pt(13)
    p2.font.color.rgb = c_muted

    bc2 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.4), Inches(5.7), Inches(4.2))
    bc2.fill.solid()
    bc2.fill.fore_color.rgb = c_dark
    bc2.line.color.rgb = c_dark
    tfc2 = bc2.text_frame
    tfc2.margin_left = Inches(0.3)
    tfc2.margin_top = Inches(0.4)
    p = tfc2.paragraphs[0]
    p.text = "LOG TEST AUTOMAT (test_critical_scenario_crash.py)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = RGBColor(56, 189, 248)
    p2 = tfc2.add_paragraph()
    p2.text = "\n[BROKER] Dispatched msg: 4f1a-b32c\n[CONSUMER] Applied local effect to DB!\n[CRASH] Process killed before ACK.\n[BROKER] ACK timeout -> Re-queueing.\n[CONSUMER-REBOOT] Re-received msg: 4f1a-b32c\n[IDEMPOTENCY] Message in processed log! Skipping effect.\n[CONSUMER] Sent ACK 4f1a-b32c\n[RESULT] 0 duplicate transactions! PASSED."
    p2.font.size = Pt(11)
    p2.font.color.rgb = RGBColor(148, 163, 184)

    # ==================== SLIDE 8: ORDERING ====================
    s8 = add_base_slide()
    add_titles(s8, "07. Păstrarea Ordinii", "Secvențiere Monotonă per Topic")

    bo1 = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.4), Inches(5.6), Inches(4.2))
    bo1.fill.solid()
    bo1.fill.fore_color.rgb = c_white
    bo1.line.color.rgb = c_border
    tfo1 = bo1.text_frame
    tfo1.margin_left = Inches(0.3)
    tfo1.margin_top = Inches(0.4)
    p = tfo1.paragraphs[0]
    p.text = "PROVOCAREA LATENȚEI ASINCRONE"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = c_primary
    p2 = tfo1.add_paragraph()
    p2.text = "\nÎn rețele distribuite, pachetele TCP pot sosi intercalate. O aplicație bancară nu poate accepta o debitare înainte de crearea contului.\n\nSistemul nominal fără număr de secvență lasă ordinea la voia sorții rețelei."
    p2.font.size = Pt(13)
    p2.font.color.rgb = c_muted

    bo2 = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.4), Inches(5.7), Inches(4.2))
    bo2.fill.solid()
    bo2.fill.fore_color.rgb = c_soft_blue
    bo2.line.color.rgb = c_primary
    tfo2 = bo2.text_frame
    tfo2.margin_left = Inches(0.3)
    tfo2.margin_top = Inches(0.4)
    p = tfo2.paragraphs[0]
    p.text = "SOLUȚIA IMPLEMENTATĂ"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = c_deep_blue
    p2 = tfo2.add_paragraph()
    p2.text = "\n• Fiecare topic deține un contor atomic monotonic protejat concurent.\n• Mesajele primesc tag-ul SequenceNumber strict crescător (1, 2, 3...).\n• Consumatorul detectează omisiuni sau reordonări și poate reordona mesajele înainte de livrarea către nucleul de business."
    p2.font.size = Pt(13)
    p2.font.color.rgb = c_dark

    # ==================== SLIDE 9: UNICAST VS MULTICAST ====================
    s9 = add_base_slide()
    add_titles(s9, "08. Rutare Flexibilă", "Topologii de Rutare: Cozi Unicast vs Topic Fan-out")

    bu1 = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.4), Inches(5.6), Inches(4.2))
    bu1.fill.solid()
    bu1.fill.fore_color.rgb = c_white
    bu1.line.color.rgb = c_border
    tfu1 = bu1.text_frame
    tfu1.margin_left = Inches(0.3)
    tfu1.margin_top = Inches(0.4)
    p = tfu1.paragraphs[0]
    p.text = "TOPIC CLASIC (MULTICAST / FAN-OUT)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = c_primary
    p2 = tfu1.add_paragraph()
    p2.text = "\n• Nume topic: news, alerts, updates\n• Toți abonații conectați primesc o copie a mesajului.\n• Replay istoric opțional la conectare.\n• Ideal pentru broadcast și invalidare de cache."
    p2.font.size = Pt(13)
    p2.font.color.rgb = c_muted

    bu2 = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.4), Inches(5.7), Inches(4.2))
    bu2.fill.solid()
    bu2.fill.fore_color.rgb = c_primary
    bu2.line.color.rgb = c_primary
    tfu2 = bu2.text_frame
    tfu2.margin_left = Inches(0.3)
    tfu2.margin_top = Inches(0.4)
    p = tfu2.paragraphs[0]
    p.text = "COADĂ DE LUCRU (UNICAST / ROUND-ROBIN)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = c_white
    p2 = tfu2.add_paragraph()
    p2.text = "\n• Nume topic: queue:tasks, queue:orders\n• Un singur consumator activ primește sarcina.\n• Distribuție echilibrată prin algoritm Round-Robin.\n• Scalare orizontală facilă prin adăugarea de noi instanțe worker."
    p2.font.size = Pt(13)
    p2.font.color.rgb = c_white

    # ==================== SLIDE 10: SUITA DE TESTE ====================
    s10 = add_base_slide()
    add_titles(s10, "09. Verificare Riguroasă", "Rezultatele Suitei de Teste Automate")

    tests_info = [
        ("test_part1.py", "Publicare, abonare, broadcast decuplat", "PASSED"),
        ("test_xml_validation.py", "Validare XSD, respingere malformed, DLQ journaling", "PASSED"),
        ("test_ordering.py", "Atribuire secvențială monotonă per topic", "PASSED"),
        ("test_consumer_ack.py", "Confirmare ACK, timeout și retransmisie automată", "PASSED"),
        ("test_unicast_queue.py", "Round-Robin pe prefixul queue:* vs Fanout topic", "PASSED"),
        ("test_critical_scenario_crash.py", "Crash înainte de ACK + deduplicare la restart", "PASSED"),
    ]
    for i, (name, desc, res) in enumerate(tests_info):
        y = Inches(2.3) + i * Inches(0.7)
        bx = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.7), Inches(0.6))
        bx.fill.solid()
        bx.fill.fore_color.rgb = c_white
        bx.line.color.rgb = c_border
        tf = bx.text_frame
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.12)
        p = tf.paragraphs[0]
        p.text = f"{name}   |   {desc}"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = c_dark
        p.font.name = "Arial"

        # Badge
        bd = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.2), y + Inches(0.1), Inches(1.1), Inches(0.4))
        bd.fill.solid()
        bd.fill.fore_color.rgb = RGBColor(220, 252, 231)
        bd.line.fill.background()
        tf_b = bd.text_frame
        tf_b.margin_top = Inches(0.05)
        p_b = tf_b.paragraphs[0]
        p_b.text = res
        p_b.font.size = Pt(10)
        p_b.font.bold = True
        p_b.font.color.rgb = c_green
        p_b.alignment = PP_ALIGN.CENTER

    # ==================== SLIDE 11: COMPARATIE ENTERPRISE ====================
    s11 = add_base_slide()
    add_titles(s11, "10. Analiză Comparativă", "Valoare Adăugată față de Soluțiile Tradiționale")

    cols_cmp = [
        ("Abordare Naivă (UDP / Sincron)", "• Pierdere imediată de mesaje la timeout\n• Risc de blocaj în cascadă\n• Tranzacții duplicate la reluare\n• Vulnerabil la mesaje corupte", c_white, c_border, c_dark, False),
        ("Sistemul Nostru Rezilient", "• Garanție At-Least-Once + Idempotency\n• Jurnal persistent JSONL pe disc\n• Validare automată XSD + DLQ\n• Rutare Unicast (Round-Robin) & Fan-out", c_primary, c_primary, c_white, True),
        ("Sisteme Enterprise (Kafka / RabbitMQ)", "• Aceleași garanții teoretice fundamentale\n• Arhitectură proiectată didactic și auditată complet\n• Fără dependențe externe masive", c_white, c_border, c_dark, False)
    ]
    for i, (title, desc, bg_c, brd_c, txt_c, is_b) in enumerate(cols_cmp):
        c_x = Inches(0.8) + i * (col_w + Inches(0.35))
        box = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x, Inches(2.4), col_w, Inches(4.2))
        box.fill.solid()
        box.fill.fore_color.rgb = bg_c
        box.line.color.rgb = brd_c
        tf = box.text_frame
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.4)
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = txt_c
        p2 = tf.add_paragraph()
        p2.text = "\n" + desc
        p2.font.size = Pt(13)
        p2.font.color.rgb = txt_c

    # ==================== SLIDE 12: DEPLOYMENT ====================
    s12 = add_base_slide()
    add_titles(s12, "11. Operare & Reproductibilitate", "Containerizare Docker & Rulare Cross-LAN")

    bd1 = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.4), Inches(5.6), Inches(4.2))
    bd1.fill.solid()
    bd1.fill.fore_color.rgb = c_white
    bd1.line.color.rgb = c_border
    tfd1 = bd1.text_frame
    tfd1.margin_left = Inches(0.3)
    tfd1.margin_top = Inches(0.4)
    p = tfd1.paragraphs[0]
    p.text = "CONTAINERIZARE DOCKER (.NET 8)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = c_primary
    p2 = tfd1.add_paragraph()
    p2.text = "\n• Imagine bazată pe .NET 8 Alpine optimizată (dimensiune redusă).\n• Expune portul 9000 pe toate interfețele (0.0.0.0).\n• Volum montat pentru persistența jurnalului de mesaje.\n• Comandă unică de pornire:\n  docker-compose up -d --build"
    p2.font.size = Pt(13)
    p2.font.color.rgb = c_muted

    bd2 = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.4), Inches(5.7), Inches(4.2))
    bd2.fill.solid()
    bd2.fill.fore_color.rgb = c_soft_blue
    bd2.line.color.rgb = c_primary
    tfd2 = bd2.text_frame
    tfd2.margin_left = Inches(0.3)
    tfd2.margin_top = Inches(0.4)
    p = tfd2.paragraphs[0]
    p.text = "COLABORARE MULTI-HOST ÎN REȚEAUA LAN"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = c_deep_blue
    p2 = tfd2.add_paragraph()
    p2.text = "\n• Fiecare membru al echipei rulează componente pe propriul laptop.\n• Publisher pe mașina 1 -> Broker pe mașina 2 -> Subscriber pe mașina 3.\n• Verifică comportamentul în condiții reale de rețea Wi-Fi cu latență fizică."
    p2.font.size = Pt(13)
    p2.font.color.rgb = c_dark

    # ==================== SLIDE 13: CONCLUZII ====================
    s13 = add_base_slide(step_text="FINAL")
    add_titles(s13, "12. Concluzii & Perspective", "Concluzii și Tranziție Către Etapa 2")

    conc = [
        ("01", "Obiectiv Atins", "Sistemul este complet decuplat temporal și spațial. Garantează fiabilitatea pe canale nesigure fără pierderi de date."),
        ("02", "Reziliență Validată", "Scenariul critic de crash post-procesare a confirmat că deduplicarea la nivel de consumator elimină complet efectele secundare."),
        ("03", "Pregătit pentru Etapa 2", "Baza de cod este modulară și pregătită pentru clustering, consens distribuit (Raft) și partiționare activă.")
    ]
    for i, (num, title, body) in enumerate(conc):
        c_x = Inches(0.8) + i * (col_w + Inches(0.35))
        box = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x, Inches(2.4), col_w, Inches(4.2))
        box.fill.solid()
        box.fill.fore_color.rgb = c_primary if i == 1 else c_white
        box.line.color.rgb = c_primary if i == 1 else c_border
        tf = box.text_frame
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.4)
        p1 = tf.paragraphs[0]
        p1.text = num
        p1.font.size = Pt(24)
        p1.font.bold = True
        p1.font.color.rgb = c_white if i == 1 else c_primary
        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(18)
        p2.font.bold = True
        p2.font.color.rgb = c_white if i == 1 else c_dark
        p3 = tf.add_paragraph()
        p3.text = "\n" + body
        p3.font.size = Pt(13)
        p3.font.color.rgb = c_white if i == 1 else c_muted

    output_path = "c:\\Users\\grigo\\Desktop\\Agent-de-mesagerie\\Prezentare_Licenta_Agent_Mesagerie.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    build_presentation()
